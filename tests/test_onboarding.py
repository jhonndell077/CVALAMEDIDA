import io
import json
import copy
import unittest
from unittest.mock import patch,MagicMock
from PIL import Image
import app
import onboarding
import user_accounts as users
from drive_bank import extract_facts,clean_facts,merge_profile
from test_telegram import MemoryStore

IDENTITY={'uid':'new-account','email':'new@example.com','name':'Alias'}
RECORD={'bank':{'folder_id':'own-bank'},'sections':{name:'own-'+str(i) for i,name in enumerate(users.SECTIONS)}}

class OnboardingTests(unittest.TestCase):
    def test_reference_cv_is_validated_and_saved_with_form_data(self):
        from werkzeug.datastructures import FileStorage
        from docx import Document
        document=Document();document.add_paragraph('Experiencia: Camarero en Restaurante X')
        stream=io.BytesIO();document.save(stream);raw=stream.getvalue()
        upload=lambda:FileStorage(stream=io.BytesIO(raw),filename='mi-cv.docx')
        files=onboarding.collect(app,{'name':'Nombre legal'},reference=upload())
        self.assertEqual(len(files),2)
        section,name,mime,data=files[-1]
        self.assertEqual(section,'Informacion personal')
        self.assertTrue(name.endswith('.docx'))
        self.assertEqual(data,raw)
        self.assertEqual(onboarding.collect(app,{},reference=upload())[0][1],name)
        store=MemoryStore();identity=users.CURRENT.set({**IDENTITY,'record':RECORD,'drive_token':'own-token'})
        try:
            with patch.object(users,'assert_drive_identity'),patch.object(onboarding,'put',return_value='saved') as put:
                result=onboarding.save(app,{},None,store,reference=upload())
                self.assertTrue(result['sync_required'])
                self.assertEqual(put.call_args.args[2],RECORD['sections']['Informacion personal'])
            for filename,data in [('cv.exe',b'bad'),('cv.pdf',b'invalid'),('cv.docx',b'invalid')]:
                with patch.object(users,'assert_drive_identity'),patch.object(onboarding,'put') as put:
                    with self.assertRaises(app.AppError):
                        onboarding.save(app,{'name':'Name'},None,store,reference=FileStorage(stream=io.BytesIO(data),filename=filename))
                    put.assert_not_called()
        finally:users.CURRENT.reset(identity)

    def test_reference_pdf_uses_deployed_reader_and_preserves_bytes(self):
        import pymupdf
        from werkzeug.datastructures import FileStorage
        with pymupdf.open() as document:
            document.new_page().insert_text((72,72),'Experiencia: Camarero')
            raw=document.tobytes()
        files=onboarding.collect(app,{'name':'Nombre legal'}, reference=FileStorage(stream=io.BytesIO(raw),filename='CV.PDF'))
        self.assertEqual(files[-1][2],'application/pdf')
        self.assertEqual(files[-1][3],raw)
        self.assertEqual(len(files),2)

    def test_password_protected_reference_does_not_write_form_or_mark_complete(self):
        import pymupdf
        from werkzeug.datastructures import FileStorage
        with pymupdf.open() as document:
            document.new_page()
            raw=document.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='password')
        store=MemoryStore();identity=users.CURRENT.set({**IDENTITY,'record':RECORD,'drive_token':'own-token'})
        try:
            with patch.object(users,'assert_drive_identity'),patch.object(onboarding,'put') as put:
                with self.assertRaises(app.AppError):
                    onboarding.save(app,{'name':'Name'},None,store,reference=FileStorage(stream=io.BytesIO(raw),filename='CV.pdf'))
                put.assert_not_called()
                self.assertIsNone(store.read(onboarding.marker(IDENTITY['uid']))[0])
        finally: users.CURRENT.reset(identity)

    def test_retry_of_saved_photo_preserves_photo_and_allows_form_to_continue(self):
        output=io.BytesIO();Image.new('RGB',(3,4),'green').save(output,format='PNG');raw=output.getvalue()
        normalized=next(item[3] for item in onboarding.collect(app,{},photo=io.BytesIO(raw)) if item[1]=='hoja_de_vida.jpeg')
        bank=MagicMock();bank.inventory.return_value={'photo':{'id':'photo','name':'hoja_de_vida.jpeg','mimeType':'image/jpeg'}}
        bank.download.return_value=(normalized,'hoja_de_vida.jpeg','image/jpeg')
        store=MemoryStore();identity=users.CURRENT.set({**IDENTITY,'record':RECORD,'drive_token':'own-token'})
        try:
            with patch.object(users,'assert_drive_identity'),patch.object(app,'BANK_CLIENT',bank),patch.object(onboarding,'put',return_value='saved') as put:
                result=onboarding.save(app,{'name':'Name'},io.BytesIO(raw),store)
                self.assertEqual(result['saved'],1)
                self.assertTrue(store.read(onboarding.marker(IDENTITY['uid']))[0]['completed'])
                self.assertEqual(put.call_args.args[3],'Datos personales - formulario.json')
                different=io.BytesIO();Image.new('RGB',(3,4),'red').save(different,format='PNG')
                with self.assertRaises(app.AppError): onboarding.save(app,{'name':'Name'},io.BytesIO(different.getvalue()),store)
                self.assertEqual(put.call_count,1)
        finally:users.CURRENT.reset(identity)

    def test_all_fields_optional_and_files_go_to_respective_sections(self):
        self.assertEqual(onboarding.collect(app,{}),[])
        files=onboarding.collect(app,{'name':'Nombre legal','education':'Bachiller','courses':'Curso de cocina',
            'experiences':'Fui camarero en X','projects':'Mi proyecto','references':'Referencia, vecino, 8090000000','skills':'Organizacion,Trabajo en equipo'})
        mapping={name:section for section,name,mime,raw in files}
        self.assertEqual(mapping['Datos personales - formulario.json'],'Informacion personal')
        self.assertEqual(mapping['Experiencias laborales - formulario.txt'],'Informacion Laboral')
        self.assertEqual(mapping['Cursos realizados - formulario.txt'],'Estudios')
        self.assertEqual(mapping['Referencias personales - formulario.txt'],'Referencias')
        self.assertEqual(mapping['Proyectos - formulario.txt'],'Conocimiento o aprendizaje')
        self.assertEqual(len(files),7)

    def test_invalid_field_or_photo_prevents_any_drive_write(self):
        with self.assertRaises(app.AppError):onboarding.collect(app,{'education':['invalid']})
        with self.assertRaises(app.AppError):onboarding.collect(app,{'name':'Name'},io.BytesIO(b'not an image'))
        image=io.BytesIO();Image.new('RGB',(4,5)).save(image,format='PNG');image.seek(0)
        files=onboarding.collect(app,{},image)
        self.assertEqual(files[0][:3],('Informacion personal','hoja_de_vida.jpeg','image/jpeg'))
        self.assertTrue(files[0][3].startswith(b'\xff\xd8'))

    def test_save_uses_own_sections_and_completion_is_only_after_all_writes(self):
        store=MemoryStore();marker=users.CURRENT.set({**IDENTITY,'record':RECORD,'drive_token':'own-token'})
        try:
            with patch.object(users,'assert_drive_identity'),patch.object(onboarding,'put',return_value='saved') as upload:
                result=onboarding.save(app,{'experiences':'Camarero','references':'Vecino'},None,store)
            self.assertEqual(result['saved'],2)
            self.assertEqual(upload.call_args_list[0].args[2],RECORD['sections']['Informacion Laboral'])
            self.assertEqual(upload.call_args_list[1].args[2],RECORD['sections']['Referencias'])
            self.assertTrue(store.read(onboarding.marker(IDENTITY['uid']))[0]['completed'])
            self.assertFalse(store.read(onboarding.marker('different-user'))[0])
            store=MemoryStore()
            with patch.object(users,'assert_drive_identity'),patch.object(onboarding,'put',side_effect=app.AppError('Drive failure')):
                with self.assertRaises(app.AppError):onboarding.save(app,{'education':'Bachiller'},None,store)
            self.assertFalse(store.read(onboarding.marker(IDENTITY['uid']))[0])
        finally:users.CURRENT.reset(marker)

    def test_status_waits_for_existing_documents_and_does_not_interrupt_existing_profiles(self):
        store=MemoryStore();core=MagicMock();core.MASTER_PROFILE=users.blank_profile(IDENTITY)
        core.load_drive_bank.return_value={'pending':0,'photo':None}
        marker=users.CURRENT.set({**IDENTITY,'record':RECORD})
        try:
            self.assertTrue(onboarding.status(core,store)['needed'])
            core.load_drive_bank.return_value={'pending':1}
            self.assertFalse(onboarding.status(core,store)['needed'])
            core.load_drive_bank.return_value={'pending':0}
            core.MASTER_PROFILE['education']=[{'title':'Bachiller'}]
            self.assertFalse(onboarding.status(core,store)['needed'])
        finally:users.CURRENT.reset(marker)
        marker=users.CURRENT.set({**IDENTITY,'record':None})
        try:
            store.write(onboarding.marker(IDENTITY['uid']),{'completed':True},0)
            self.assertFalse(onboarding.status(core,store)['needed'])
        finally:users.CURRENT.reset(marker)

    def test_structured_personal_file_reads_literal_fields_without_ai_and_deletion_removes_them(self):
        raw=json.dumps({'cv_bank_form':1,'candidate':{'name':'Nombre legal','address':'Mi direccion'}}).encode()
        facts=extract_facts('Datos personales - formulario.json','application/json',raw,'Alias')
        base=users.blank_profile(IDENTITY)
        docs={'personal':{'active':True,'status':'incorporado','name':'Datos personales','fingerprint':'v1','facts':facts},
              'job':{'active':True,'status':'incorporado','name':'Trabajo','fingerprint':'v1','facts':clean_facts({'recipient':'Nombre legal','experiences':[{'title':'Camarero','employer':'X'}],'generated_skills':['Atencion al cliente']})}}
        profile,_=merge_profile(base,docs)
        self.assertEqual(profile['candidate']['name'],'Nombre legal')
        self.assertEqual(len(profile['experiences']),1)
        self.assertEqual(profile['skills'][0]['origin'],'sugerida_por_experiencia_o_estudios')
        profile,_=merge_profile(base,{})
        self.assertEqual(profile['candidate']['address'],'')
        self.assertEqual(profile['skills'],[])

    def test_explicit_skills_replace_suggestions_and_no_evidence_generates_no_skills(self):
        base=users.blank_profile(IDENTITY)
        doc={'active':True,'status':'incorporado','name':'Trabajo','fingerprint':'v1','facts':clean_facts({'experiences':[{'title':'Camarero'}],'skills':['Servicio'],'generated_skills':['Sugerida']})}
        profile,_=merge_profile(base,{'doc':doc})
        self.assertEqual([s['label'] for s in profile['skills']],['Servicio'])
        doc['facts']['skills']=[];doc['facts']['experiences']=[]
        profile,_=merge_profile(base,{'doc':doc})
        self.assertEqual(profile['skills'],[])

    def test_registered_user_can_continue_with_academic_data_and_no_employment(self):
        profile=users.blank_profile(IDENTITY)
        profile['education']=[{'id':'study-one','title':'Bachiller','institution':'Escuela'}]
        marker=users.CURRENT.set(IDENTITY)
        try:
            with patch.object(app,'MASTER_PROFILE',profile),patch.object(app,'EXPERIENCE_BY_ID',{}),patch.object(app,'SKILL_BY_ID',{}),patch.object(app,'COURSE_BY_ID',{}):
                analysis=app.sanitize_analysis({'vacancy':{'position':'Asistente'},'matching':{}})
                draft=app.make_draft(analysis)
                self.assertEqual(draft['experiences'],[])
                self.assertEqual(draft['education'],profile['education'])
                self.assertIn('Bachiller',draft['summary'])
                self.assertNotIn('proyectos personales',draft['summary'])
                with patch.object(app,'current_photo_path',return_value=__import__('pathlib').Path('tmp/no-onboarding-photo')):
                    self.assertTrue(app.make_pdf(draft).getvalue().startswith(b'%PDF'))
        finally:users.CURRENT.reset(marker)


class ReferenceImportTests(unittest.TestCase):
    def test_import_returns_editable_explicit_fields_without_generated_claims(self):
        from werkzeug.datastructures import FileStorage
        from docx import Document
        document = Document(); document.add_paragraph('Nombre y experiencia de prueba')
        stream = io.BytesIO(); document.save(stream); stream.seek(0)
        facts = {'candidate': {'name': 'Ana', 'phones': ['123'], 'residence': 'Santiago'},
                 'experiences': [{'title': 'Auxiliar', 'employer': 'Empresa', 'responsibilities': ['Atender'],
                                  'generated_responsibilities': ['Inventada'], 'id': 'source'}],
                 'education': [{'title': 'Bachiller'}], 'skills': ['Comunicación'],
                 'generated_skills': ['Sugerida'], 'notes': ['Fecha ilegible']}
        with patch('drive_bank.extract_facts', return_value=facts):
            result = onboarding.import_reference(app, FileStorage(stream=stream, filename='cv.docx'))
        self.assertEqual(result['fields']['name'], 'Ana')
        self.assertEqual(result['fields']['address'], 'Santiago')
        self.assertEqual(result['fields']['experiences'], [{'title':'Auxiliar','employer':'Empresa','responsibilities':['Atender']}])
        self.assertEqual(result['fields']['skills'], ['Comunicación'])
        self.assertEqual(result['notes'], ['Fecha ilegible'])
        onboarding.collect(app, result['fields'])

    def test_missing_reference_rejected(self):
        with self.assertRaises(app.AppError): onboarding.import_reference(app, None)
