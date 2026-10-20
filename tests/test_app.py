import base64
import copy
import io
import json
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

import pymupdf as fitz
from docx import Document
from PIL import Image


APP_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP_DIR))
import app as cv_app  # noqa: E402


def valid_analysis():
    return {
        "vacancy": {"position": "Auxiliar de almacén"},
        "matching": {
            "template": "supermercados",
            "selected_experiences": [
                {
                    "experience_id": "caribe-gondolero",
                    "responsibility_ids": ["reponer-organizar", "asistencia-productos"],
                },
                {
                    "experience_id": "liverpool-gondolero",
                    "responsibility_ids": ["colocacion-rotacion", "presentacion-disponibilidad"],
                },
                {
                    "experience_id": "grupo-ramos-almacen",
                    "responsibility_ids": ["recepcion-clasificacion"],
                },
            ],
            "selected_skills": ["inventario", "recepcion-almacen", {"invalid": True}],
            "selected_course_ids": [
                "paquete-oficina",
                "manejo-alimentos",
                "invalid-course",
            ],
            "selected_project_ids": ["gestor-inventario-wtf", "invalid-id"],
            "items": [],
            "gaps": [],
        },
    }


def supervisor_analysis():
    return {
        "vacancy": {"position": "Supervisor de Restaurante"},
        "matching": {
            "template": "oficina",
            "selected_experiences": [
                {
                    "experience_id": "farah-chef",
                    "responsibility_ids": [
                        "admin-general",
                        "control-costos",
                        "planificacion-procesos",
                    ],
                },
                {
                    "experience_id": "farah-sous-chef",
                    "responsibility_ids": [
                        "coordina-produccion",
                        "supervisa-cocina",
                        "cumple-recetas-calidad",
                        "gestiona-inventario-cocina",
                    ],
                },
                {
                    "experience_id": "reysa-crew-leader",
                    "responsibility_ids": [
                        "seguimiento-recetas",
                        "trabajo-organizado",
                    ],
                },
                {
                    "experience_id": "cris-asistente-admin",
                    "responsibility_ids": ["coordina-internas", "manejo-recursos"],
                },
                {
                    "experience_id": "grupo-ramos-cajero",
                    "responsibility_ids": ["servicio-rapido"],
                },
            ],
            "selected_skills": [
                "supervision",
                "coordinacion",
                "procedimientos-estandares",
                "calidad-higiene",
                "costos-planificacion",
                "inventario",
            ],
            "selected_course_ids": [
                "manejo-alimentos",
                "pastas-artesanales",
                "seguridad-higiene",
            ],
            "selected_project_ids": [],
        },
    }


class ResumeAppTests(unittest.TestCase):
    def setUp(self):
        archive_env = patch.dict("os.environ", {"CV_ARCHIVE_PROVIDER": "legacy"})
        archive_env.start()
        self.addCleanup(archive_env.stop)

    def test_drive_export_uploads_and_confirms_fixed_parent(self):
        folder = {"id": "13LPPV7i7FycIGAqjokJPKsfbt2h6Ahk9", "name": "CV GENERADOS"}
        with (
            patch.dict("os.environ", {"CV_ARCHIVE_PROVIDER": "drive"}),
            patch("app.verify_drive_access", return_value=folder),
            patch("app.drive_request", side_effect=[{"id": "saved-file"}, {"id": "saved-file", "parents": [folder["id"]], "webViewLink": "https://drive.google.com/file/d/saved-file/view"}]) as upload,
        ):
            response = cv_app.app.test_client().post("/api/export/pdf", json={"analysis": valid_analysis()}, headers={"X-Drive-Access-Token": "test-token"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["X-CV-Drive-URL"], "https://drive.google.com/file/d/saved-file/view")
        body = upload.call_args_list[0].args[2]
        self.assertIn(folder["id"].encode(), body)
        self.assertIn(b"%PDF", body)
        self.assertEqual(upload.call_count, 2)

    def test_drive_export_rejects_missing_authorization(self):
        with patch.dict("os.environ", {"CV_ARCHIVE_PROVIDER": "drive"}):
            response = cv_app.app.test_client().post("/api/export/pdf", json={"analysis": valid_analysis()})
        self.assertEqual(response.status_code, 401)

    def test_drive_rejects_different_account(self):
        with patch("app.drive_request", return_value={"email": "someone@example.com", "email_verified": True}) as api:
            with self.assertRaises(cv_app.AppError) as error:
                cv_app.verify_drive_access("test-token")
        self.assertEqual(error.exception.status_code, 403)
        self.assertEqual(api.call_count, 1)

    def test_drive_rejects_wrong_upload_parent(self):
        with (
            patch.dict("os.environ", {"CV_ARCHIVE_PROVIDER": "drive"}),
            cv_app.app.test_request_context(headers={"X-Drive-Access-Token": "test-token"}),
            patch("app.verify_drive_access", return_value={"id": "target"}),
            patch("app.drive_request", side_effect=[{"id": "saved"}, {"parents": ["other"]}]),
        ):
            with self.assertRaises(cv_app.AppError):
                cv_app.archive_generated_resume("cv.pdf", "pdf", b"%PDF")

    def test_ollama_stream_response_is_reassembled(self):
        response = MagicMock()
        response.__iter__.return_value = iter(
            [
                b'{"message":{"content":"{\\"position\\":"},"done":false}\n',
                b'{"message":{"content":"\\"cajero\\"}"},"done":true}\n',
            ]
        )
        request_context = MagicMock()
        request_context.__enter__.return_value = response
        with patch("app.urllib.request.urlopen", return_value=request_context):
            result = cv_app.ollama_json(
                "/api/chat",
                {"stream": True},
            )
        self.assertEqual(
            result["message"]["content"], '{"position":"cajero"}'
        )

    def test_analysis_uses_local_whitelisted_profile_facts(self):
        model_response = {
            "message": {
                "content": json.dumps(
                    {
                        "vacancy": {"position": "Auxiliar de almacén"},
                        "matching": {
                            "items": [],
                            "selected_experiences": [
                                {
                                    "experience_id": "grupo-ramos-almacen",
                                    "responsibility_ids": ["recepcion-clasificacion"],
                                }
                            ],
                            "selected_skills": ["recepcion-almacen"],
                            "selected_course_ids": ["paquete-oficina"],
                            "selected_project_ids": [],
                            "gaps": [],
                            "template": "supermercados",
                            "template_reason": "Operaciones de retail",
                        },
                    }
                )
            }
        }
        with (
            patch("app.require_model", return_value="llama3.1:8b"),
            patch("app.ollama_json", return_value=model_response) as model_call,
        ):
            result = cv_app.analyze_with_model(
                "llama3.1:8b", "Auxiliar de almacén", []
            )

        self.assertEqual(
            result["matching"]["selected_experiences"][0]["experience_id"],
            "grupo-ramos-almacen",
        )
        payload = model_call.call_args.args[1]
        self.assertIsInstance(payload["format"], dict)
        user_content = payload["messages"][1]["content"]
        system_content = payload["messages"][0]["content"]
        self.assertIn("grupo-ramos-almacen", user_content)
        self.assertIn("Manejador de Paquete de Oficina", user_content)
        self.assertIn("no copies la lista completa de cursos", system_content)
        self.assertEqual(
            result["matching"]["selected_course_ids"], ["paquete-oficina"]
        )
        self.assertNotIn(cv_app.MASTER_PROFILE["candidate"]["email"], user_content)
        self.assertNotIn(cv_app.MASTER_PROFILE["candidate"]["national_id"], user_content)

    def test_vertex_analysis_uses_google_model_without_contact_details(self):
        model_result = {
            "vacancy": {"position": "Auxiliar de almacén"},
            "matching": {
                "items": [],
                "selected_experiences": [
                    {
                        "experience_id": "grupo-ramos-almacen",
                        "responsibility_ids": ["recepcion-clasificacion"],
                    }
                ],
                "selected_skills": ["recepcion-almacen"],
                "selected_course_ids": ["paquete-oficina"],
                "selected_project_ids": [],
                "gaps": [],
                "template": "supermercados",
                "template_reason": "Operaciones de retail",
            },
        }
        with (
            patch.dict(
                "os.environ",
                {
                    "CV_AI_PROVIDER": "vertex",
                    "GOOGLE_CLOUD_PROJECT": "cv-a-la-medida-jhonn",
                },
            ),
            patch("app.vertex_json", return_value=json.dumps(model_result)) as model_call,
        ):
            result = cv_app.analyze_with_model(
                cv_app.VERTEX_MODEL, "Auxiliar de almacén", []
            )

        self.assertEqual(
            result["matching"]["selected_experiences"][0]["experience_id"],
            "grupo-ramos-almacen",
        )
        self.assertEqual(model_call.call_args.args[0], cv_app.VERTEX_MODEL)
        prompt = model_call.call_args.args[1]
        self.assertIn("grupo-ramos-almacen", prompt)
        self.assertNotIn(cv_app.MASTER_PROFILE["candidate"]["email"], prompt)
        self.assertNotIn(cv_app.MASTER_PROFILE["candidate"]["national_id"], prompt)

    def test_vertex_model_requires_project_and_rejects_other_models(self):
        with patch.dict("os.environ", {"CV_AI_PROVIDER": "vertex"}, clear=False):
            with patch.dict("os.environ", {"GOOGLE_CLOUD_PROJECT": ""}):
                with self.assertRaises(cv_app.AppError) as missing_project:
                    cv_app.require_model(cv_app.VERTEX_MODEL)
            self.assertEqual(missing_project.exception.status_code, 503)
            with patch.dict(
                "os.environ",
                {"GOOGLE_CLOUD_PROJECT": "cv-a-la-medida-jhonn"},
            ):
                with self.assertRaises(cv_app.AppError):
                    cv_app.require_model("unapproved-model")

    def test_vertex_sdk_builds_structured_request(self):
        from google.genai import types

        client = MagicMock()
        client.models.generate_content.return_value.text = '{"ready":true}'
        schema = {"type": "OBJECT", "properties": {"ready": {"type": "BOOLEAN"}}}
        image = base64.b64encode(b"image-bytes").decode("ascii")
        with (
            patch.dict(
                "os.environ",
                {
                    "GOOGLE_CLOUD_PROJECT": "cv-a-la-medida-jhonn",
                    "GOOGLE_CLOUD_LOCATION": "us-central1",
                },
            ),
            patch("google.genai.Client", return_value=client) as client_factory,
        ):
            result = cv_app.vertex_json(
                cv_app.VERTEX_MODEL, "Analiza la vacante", [image], schema
            )

        self.assertEqual(result, '{"ready":true}')
        client_factory.assert_called_once_with(
            vertexai=True,
            project="cv-a-la-medida-jhonn",
            location="us-central1",
            http_options=client_factory.call_args.kwargs["http_options"],
        )
        self.assertEqual(client_factory.call_args.kwargs["http_options"].retry_options.attempts, 1)
        self.assertEqual(client_factory.call_args.kwargs["http_options"].timeout, 75000)
        request = client.models.generate_content.call_args.kwargs
        self.assertEqual(request["model"], cv_app.VERTEX_MODEL)
        self.assertEqual(request["config"].response_schema, schema)
        self.assertEqual(request["config"].response_mime_type, "application/json")
        self.assertEqual(len(request["contents"]), 2)
        self.assertIsInstance(request["contents"][1], types.Part)

    def test_vertex_schema_stays_bounded_when_evidence_bank_grows(self):
        schema = cv_app.analysis_output_schema()
        original = copy.deepcopy(schema)
        matching = schema["properties"]["matching"]["properties"]
        matching["selected_skills"]["items"]["enum"] = [
            f"drive-skill-{index:016x}" for index in range(2000)
        ]
        matching["selected_experiences"]["items"]["properties"]["responsibility_ids"]["items"]["enum"] = [
            f"drive-task-{index:016x}" for index in range(2000)
        ]
        compact = cv_app.vertex_response_schema(schema)
        self.assertNotIn('maxItems', json.dumps(compact))
        self.assertLess(len(json.dumps(compact)), 5000)
        self.assertNotIn("drive-skill-", json.dumps(compact))
        self.assertNotIn("drive-task-", json.dumps(compact))
        self.assertEqual(compact["required"], schema["required"])
        self.assertEqual(
            compact["properties"]["matching"]["properties"]["template"]["enum"],
            original["properties"]["matching"]["properties"]["template"]["enum"],
        )
        self.assertEqual(len(matching["selected_skills"]["items"]["enum"]), 2000)

    def test_vacancy_file_formats_are_extracted_locally(self):
        text_file = base64.b64encode(
            "Asistente administrativo".encode("utf-8")
        ).decode("ascii")
        text, images = cv_app.extract_vacancy(
            {"file_name": "vacante.txt", "file_data": text_file}
        )
        self.assertIn("Asistente administrativo", text)
        self.assertEqual(images, [])

        word = Document()
        word.add_paragraph("Recepción de mercancía y almacenamiento")
        word_buffer = io.BytesIO()
        word.save(word_buffer)
        text, images = cv_app.extract_vacancy(
            {
                "file_name": "vacante.docx",
                "file_data": base64.b64encode(word_buffer.getvalue()).decode("ascii"),
            }
        )
        self.assertIn("almacenamiento", text)
        self.assertEqual(images, [])

        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((40, 50), "Puesto auxiliar con atención al cliente y experiencia requerida")
        text, images = cv_app.extract_vacancy(
            {
                "file_name": "vacante.pdf",
                "file_data": base64.b64encode(pdf.tobytes()).decode("ascii"),
            }
        )
        self.assertIn("atención al cliente", text)
        self.assertEqual(images, [])

        image_buffer = io.BytesIO()
        Image.new("RGB", (80, 40), "white").save(image_buffer, format="PNG")
        text, images = cv_app.extract_vacancy(
            {
                "file_name": "anuncio.png",
                "file_data": base64.b64encode(image_buffer.getvalue()).decode("ascii"),
            }
        )
        self.assertEqual(text, "")
        self.assertEqual(len(images), 1)
        self.assertTrue(base64.b64decode(images[0]).startswith(b"\xff\xd8"))

        scan = fitz.open()
        scan_page = scan.new_page()
        scan_page.insert_image(scan_page.rect, stream=image_buffer.getvalue())
        text, images = cv_app.extract_vacancy(
            {
                "file_name": "anuncio-escaneado.pdf",
                "file_data": base64.b64encode(scan.tobytes()).decode("ascii"),
            }
        )
        self.assertEqual(text, "")
        self.assertEqual(len(images), 1)

    def test_master_profile_has_photo_and_preserved_information(self):
        candidate = cv_app.MASTER_PROFILE["candidate"]
        self.assertTrue((APP_DIR / candidate["photo"]).is_file())
        self.assertEqual(candidate["photo"], "data/hoja_de_vida.jpeg")
        self.assertEqual(candidate["name"], "Jhonn Angel Santana Castillo")
        self.assertTrue(candidate["national_id"])
        self.assertEqual(len(candidate["phones"]), 3)
        self.assertTrue(candidate["email"])
        self.assertTrue(candidate["address"])
        self.assertTrue(candidate["residence"])
        self.assertGreaterEqual(len(cv_app.MASTER_PROFILE["education"]), 2)
        self.assertGreaterEqual(len(cv_app.MASTER_PROFILE["courses"]), 1)
        self.assertGreaterEqual(len(cv_app.MASTER_PROFILE["references"]), 1)

    def test_cloud_archive_uploads_to_private_storage_bucket(self):
        storage_client = MagicMock()
        with (
            patch.dict(
                "os.environ",
                {
                    "CV_ALLOW_REMOTE": "1",
                    "CV_RESUME_ARCHIVE_BUCKET": "cv-archive.example",
                },
            ),
            patch("google.cloud.storage.Client", return_value=storage_client),
        ):
            location = cv_app.archive_generated_resume(
                "Jhonn_Santana_Supervisor_20261006_120000_ab12cd34.pdf",
                "pdf",
                b"%PDF-test",
            )

        self.assertEqual(
            location,
            "gs://cv-archive.example/CV generados/2026/10/"
            "Jhonn_Santana_Supervisor_20261006_120000_ab12cd34.pdf",
        )
        storage_client.bucket.assert_called_once_with("cv-archive.example")
        storage_client.bucket.return_value.blob.return_value.upload_from_string.assert_called_once_with(
            b"%PDF-test",
            content_type="application/pdf",
        )

    def test_cloud_archive_requires_a_configured_bucket(self):
        with patch.dict(
            "os.environ",
            {"CV_ALLOW_REMOTE": "1", "CV_RESUME_ARCHIVE_BUCKET": ""},
            clear=False,
        ):
            with patch.dict("os.environ", {"FIREBASE_CONFIG": ""}):
                with self.assertRaises(cv_app.AppError) as error:
                    cv_app.archive_generated_resume("resume.pdf", "pdf", b"%PDF")
        self.assertEqual(error.exception.status_code, 503)
        self.assertIn("CV_RESUME_ARCHIVE_BUCKET", str(error.exception))

    def test_draft_keeps_duplicate_roles_as_separate_employments(self):
        draft = cv_app.make_draft(valid_analysis())
        self.assertEqual(draft["template"], "supermercados")
        gondolero = [item for item in draft["experiences"] if item["title"] == "Gondolero"]
        self.assertEqual(len(gondolero), 2)
        self.assertEqual([item["employer"] for item in gondolero],
                         ["Supermercado Caribe", "Liverpool Supermercados"])
        for item in gondolero:
            self.assertIsInstance(item["id"], str)
            self.assertEqual(len(item["selected_responsibilities"]), 2)
            self.assertTrue(all(isinstance(text, str) for text in item["selected_responsibilities"]))
        self.assertEqual(
            [item["id"] for item in draft["courses"]], ["paquete-oficina"]
        )
        self.assertEqual(
            len(draft["references"]), len(cv_app.MASTER_PROFILE["references"])
        )
        self.assertEqual(draft["education"], cv_app.MASTER_PROFILE["education"])
        self.assertEqual(
            [item["id"] for item in draft["projects"]], ["gestor-inventario-wtf"]
        )
        self.assertNotIn({"invalid": True}, draft["skills"])

    def test_office_vacancy_excludes_food_courses(self):
        analysis = valid_analysis()
        analysis["vacancy"]["position"] = "Asistente Administrativo"
        analysis["matching"]["selected_experiences"] = [
            {
                "experience_id": "cris-asistente-admin",
                "responsibility_ids": ["gestion-documentacion"],
            }
        ]
        draft = cv_app.make_draft(analysis)
        self.assertEqual(
            [item["id"] for item in draft["courses"]], ["paquete-oficina"]
        )

    def test_invalid_experience_ids_are_rejected(self):
        analysis = {
            "vacancy": {"position": "Recepcionista"},
            "matching": {
                "selected_experiences": [
                    {"experience_id": ["invalid"], "responsibility_ids": []}
                ]
            },
        }
        with self.assertRaises(cv_app.AppError):
            cv_app.make_draft(analysis)

    def test_pdf_and_docx_exports_keep_accents_and_fixed_sections(self):
        draft = cv_app.make_draft(valid_analysis())
        self.assertEqual(draft["template"], "supermercados")
        pdf = cv_app.make_pdf(draft)
        pdf_document = fitz.open(stream=pdf.getvalue(), filetype="pdf")
        photo_rect = pdf_document[0].get_image_rects(pdf_document[0].get_images()[0][0])[0]
        sidebar_rgb = tuple(value / 255 for value in cv_app.rgb_from_hex(cv_app.TEMPLATE_SIDEBARS[draft["template"]]))
        sidebar_rect = next(drawing["rect"] for drawing in pdf_document[0].get_drawings()
            if drawing.get("fill") and all(abs(a-b) < 0.00001 for a, b in zip(drawing["fill"], sidebar_rgb)))
        self.assertAlmostEqual((photo_rect.x0 + photo_rect.x1) / 2,
            (sidebar_rect.x0 + sidebar_rect.x1) / 2, places=3)
        pdf_text = re.sub(
            r"\s+",
            " ",
            "\n".join(page.get_text() for page in pdf_document),
        )
        self.assertIn("Jhonn Angel Santana Castillo", pdf_text)
        self.assertIn("formación académica", pdf_text.casefold())
        self.assertIn("cursos realizados", pdf_text.casefold())
        self.assertIn("referencias personales", pdf_text.casefold())
        self.assertIn("Liverpool Supermercados", pdf_text)
        self.assertIn("Manejador de Paquete de Oficina", pdf_text)
        self.assertNotIn("Manejo de alimentos", pdf_text)
        self.assertNotIn("Pastas Artesanales", pdf_text)
        for item in draft["courses"]:
            self.assertIn(item["title"], pdf_text)
        for item in draft["references"]:
            self.assertIn(item["name"], pdf_text)
        self.assertEqual(pdf_document.page_count, 1)
        section_positions = [
            pdf_text.casefold().index(section)
            for section in (
                "sobre mi",
                "formación académica",
                "experiencia laboral",
                "referencias personales",
            )
        ]
        self.assertEqual(section_positions, sorted(section_positions))

        docx = cv_app.make_docx(draft)
        document = Document(io.BytesIO(docx.getvalue()))
        self.assertIn(
            'w:fill="fbefe4"',
            document.tables[0].cell(0, 0)._tc.xml.casefold(),
        )
        docx_text = "\n".join(
            paragraph.text for paragraph in document.paragraphs
        ) + "\n" + "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        self.assertIn("Jhonn Angel Santana Castillo", docx_text)
        self.assertIn("formación académica", docx_text.casefold())
        self.assertIn("cursos realizados", docx_text.casefold())
        self.assertIn("referencias personales", docx_text.casefold())
        self.assertIn("Liverpool Supermercados", docx_text)
        self.assertIn("Manejador de Paquete de Oficina", docx_text)
        self.assertNotIn("Manejo de alimentos", docx_text)
        self.assertNotIn("Pastas Artesanales", docx_text)
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertEqual(len(document.tables), 1)
        references_tables = document.tables[0].cell(0, 1).tables
        reference_names = " ".join(
            cell.text
            for table in references_tables
            for row in table.rows
            for cell in row.cells
        )
        for item in draft["courses"]:
            self.assertIn(item["title"], docx_text)
        for item in draft["references"]:
            self.assertIn(item["name"], reference_names)
        docx_sections = [
            docx_text.casefold().index(section)
            for section in (
                "sobre mi",
                "formación académica",
                "experiencia laboral",
                "referencias personales",
            )
        ]
        self.assertEqual(docx_sections, sorted(docx_sections))

    def test_supervisor_draft_uses_kitchen_template_and_fits_one_page(self):
        analysis = supervisor_analysis()
        analysis["matching"]["template"] = "cocina"
        draft = cv_app.make_draft(analysis)
        self.assertEqual(draft["template"], "cocina")
        self.assertEqual(
            [item["id"] for item in draft["courses"]],
            ["manejo-alimentos", "pastas-artesanales", "seguridad-higiene"],
        )
        pdf = fitz.open(
            stream=cv_app.make_pdf(draft).getvalue(),
            filetype="pdf",
        )
        self.assertEqual(pdf.page_count, 1)
        text = re.sub(r"\s+", " ", pdf[0].get_text()).casefold()
        self.assertIn("chef: farah modern kitchen", text)
        self.assertIn("supervisaba al personal de cocina", text)
        self.assertIn("referencias personales", text)
        self.assertNotIn("proyectos relevantes", text)

    def test_api_exports_and_profile_summary_are_available_locally(self):
        client = cv_app.app.test_client()
        foreign_host = client.get("/", headers={"Host": "example.com"})
        self.assertEqual(foreign_host.status_code, 403)
        profile_response = client.get("/api/profile-summary")
        self.assertEqual(profile_response.status_code, 200)
        self.assertGreater(len(profile_response.json["experiences"]), 0)
        photo_response = client.get("/api/profile-photo")
        self.assertEqual(photo_response.status_code, 200)
        self.assertEqual(photo_response.mimetype, "image/jpeg")
        photo_response.close()

        response = client.post("/api/draft", json={"analysis": valid_analysis()})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["draft"]["position"], "Auxiliar de almacén")

        with tempfile.TemporaryDirectory() as archive_dir:
            with patch.object(
                cv_app, "GENERATED_RESUMES_DIR", Path(archive_dir)
            ):
                pdf_response = client.post(
                    "/api/export/pdf", json={"analysis": valid_analysis()}
                )
                docx_response = client.post(
                    "/api/export/docx", json={"analysis": valid_analysis()}
                )
            archived_files = list(Path(archive_dir).iterdir())
        self.assertEqual(pdf_response.status_code, 200)
        self.assertEqual(docx_response.status_code, 200)
        self.assertTrue(pdf_response.data.startswith(b"%PDF"))
        self.assertTrue(docx_response.data.startswith(b"PK"))
        self.assertEqual({path.suffix for path in archived_files}, {".pdf", ".docx"})
        self.assertEqual(len(archived_files), 2)

    def test_local_profile_editor_saves_new_experiences_to_master_file(self):
        client = cv_app.app.test_client()
        original_profile = cv_app.MASTER_PROFILE
        profile = copy.deepcopy(original_profile)
        profile["experiences"].append(
            {
                "id": "experiencia-prueba",
                "title": "Asistente de prueba",
                "employer": "Empresa de prueba",
                "category": "oficina",
                "responsibilities": [
                    {"id": "funcion-prueba", "text": "Organización de documentos."}
                ],
            }
        )
        with tempfile.TemporaryDirectory() as profile_dir:
            profile_path = Path(profile_dir) / "master_profile.json"
            try:
                with (
                    patch.object(cv_app, "PROFILE_PATH", profile_path),
                    patch.dict(
                        "os.environ",
                        {"CV_ALLOW_REMOTE": "", "CV_AUTH_REQUIRED": ""},
                    ),
                ):
                    response = client.put("/api/profile", json=profile)
                    self.assertEqual(response.status_code, 200, response.json)
                    self.assertEqual(
                        response.json["profile"]["experiences"][-1]["id"],
                        "experiencia-prueba",
                    )
                    self.assertEqual(
                        json.loads(profile_path.read_text(encoding="utf-8"))[
                            "experiences"
                        ][-1]["id"],
                        "experiencia-prueba",
                    )
                    self.assertIn("experiencia-prueba", cv_app.EXPERIENCE_BY_ID)
            finally:
                cv_app.MASTER_PROFILE = original_profile
                cv_app.refresh_profile_indexes()

    def test_profile_editor_rejects_invalid_data_and_cloud_writes(self):
        client = cv_app.app.test_client()
        invalid_profile = copy.deepcopy(cv_app.MASTER_PROFILE)
        invalid_profile["experiences"][0]["responsibilities"][0]["id"] = "bad id"
        with tempfile.TemporaryDirectory() as profile_dir:
            profile_path = Path(profile_dir) / "master_profile.json"
            with (
                patch.object(cv_app, "PROFILE_PATH", profile_path),
                patch.dict(
                    "os.environ",
                    {"CV_ALLOW_REMOTE": "", "CV_AUTH_REQUIRED": ""},
                ),
            ):
                invalid_response = client.put("/api/profile", json=invalid_profile)
                self.assertEqual(invalid_response.status_code, 400)
                self.assertFalse(profile_path.exists())
                empty_profile = copy.deepcopy(cv_app.MASTER_PROFILE)
                empty_profile["experiences"] = []
                empty_response = client.put("/api/profile", json=empty_profile)
                self.assertEqual(empty_response.status_code, 400)
                self.assertFalse(profile_path.exists())
            with patch.dict(
                "os.environ",
                {"CV_ALLOW_REMOTE": "1", "CV_AUTH_REQUIRED": ""},
            ):
                cloud_response = client.put(
                    "/api/profile", json=cv_app.MASTER_PROFILE
                )
            self.assertEqual(cloud_response.status_code, 403)

    def test_cloud_profile_endpoints_require_google_authentication(self):
        client = cv_app.app.test_client()
        with patch.dict(
            "os.environ",
            {
                "CV_AUTH_REQUIRED": "1",
                "CV_ALLOWED_EMAILS": "jhonndell077@gmail.com",
            },
        ):
            response = client.get("/api/profile-summary")
            full_profile_response = client.get("/api/profile")
            photo_response = client.get("/api/profile-photo")
            health = client.get("/api/health")
        self.assertEqual(response.status_code, 401)
        self.assertEqual(full_profile_response.status_code, 401)
        self.assertEqual(photo_response.status_code, 401)
        self.assertIn("Inicia sesión", response.json["error"])
        self.assertEqual(health.status_code, 200)
        self.assertTrue(health.json["auth_required"])

    def test_knowledge_map_ingestion_is_incremental_and_pending(self):
        client = cv_app.app.test_client()
        responsibility = cv_app.MASTER_PROFILE["experiences"][0][
            "responsibilities"
        ][0]
        original_profile = copy.deepcopy(cv_app.MASTER_PROFILE)
        document = f"Historial documentado: {responsibility['text']}".encode()
        with tempfile.TemporaryDirectory() as map_dir:
            map_path = Path(map_dir) / "knowledge_map.json"
            with (
                patch.object(cv_app, "KNOWLEDGE_MAP_PATH", map_path),
                patch.dict(
                    "os.environ",
                    {"CV_ALLOW_REMOTE": "", "CV_AUTH_REQUIRED": ""},
                ),
            ):
                first = client.post(
                    "/api/knowledge-map/documents",
                    data={
                        "files": (io.BytesIO(document), "historial.txt"),
                        "modifiedTimes": json.dumps(
                            ["2026-10-06T10:00:00Z"]
                        ),
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(first.status_code, 200, first.json)
                first_map = first.json["knowledge_map"]
                self.assertEqual(first.json["results"][0]["status"], "processed")
                self.assertEqual(len(first_map["documents"]), 1)
                evidence = next(iter(first_map["evidence"].values()))
                self.assertEqual(evidence["status"], "pending_validation")
                self.assertIn(
                    responsibility["id"],
                    [reference["id"] for reference in evidence["references"]],
                )
                self.assertEqual(cv_app.MASTER_PROFILE, original_profile)

                with patch(
                    "app.extract_historical_document",
                    side_effect=AssertionError("unchanged file was reprocessed"),
                ):
                    unchanged = client.post(
                        "/api/knowledge-map/documents",
                        data={
                            "files": (io.BytesIO(document), "historial.txt"),
                            "modifiedTimes": json.dumps(
                                ["2026-10-06T11:00:00Z"]
                            ),
                        },
                        content_type="multipart/form-data",
                    )
                self.assertEqual(unchanged.status_code, 200, unchanged.json)
                self.assertEqual(
                    unchanged.json["results"][0]["status"], "unchanged"
                )
                self.assertEqual(
                    unchanged.json["knowledge_map"]["documents"][
                        "local:historial.txt"
                    ]["lastProcessedAt"],
                    first_map["documents"]["local:historial.txt"][
                        "lastProcessedAt"
                    ],
                )

                changed = client.post(
                    "/api/knowledge-map/documents",
                    data={
                        "files": (
                            io.BytesIO(document + b"\nNueva anotacion"),
                            "historial.txt",
                        ),
                        "modifiedTimes": json.dumps(
                            ["2026-10-07T10:00:00Z"]
                        ),
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(changed.status_code, 200, changed.json)
                updated_document = changed.json["knowledge_map"]["documents"][
                    "local:historial.txt"
                ]
                self.assertEqual(
                    changed.json["results"][0]["status"], "updated"
                )
                self.assertEqual(updated_document["revisionCount"], 2)
                self.assertEqual(len(updated_document["revisions"]), 2)
                self.assertEqual(
                    len(changed.json["knowledge_map"]["evidence"]), 2
                )
                self.assertEqual(cv_app.MASTER_PROFILE, original_profile)

    def test_knowledge_map_rejects_cloud_writes_and_unsupported_documents(self):
        client = cv_app.app.test_client()
        with tempfile.TemporaryDirectory() as map_dir:
            with (
                patch.object(
                    cv_app,
                    "KNOWLEDGE_MAP_PATH",
                    Path(map_dir) / "knowledge.json",
                ),
                patch.dict(
                    "os.environ",
                    {"CV_ALLOW_REMOTE": "1", "CV_AUTH_REQUIRED": ""},
                ),
            ):
                remote_response = client.post(
                    "/api/knowledge-map/documents",
                    data={
                        "files": (io.BytesIO(b"contenido"), "historial.txt"),
                        "modifiedTimes": json.dumps([None]),
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(remote_response.status_code, 403)
            with patch.dict(
                "os.environ",
                {"CV_ALLOW_REMOTE": "", "CV_AUTH_REQUIRED": ""},
            ):
                unsupported_response = client.post(
                    "/api/knowledge-map/documents",
                    data={
                        "files": (io.BytesIO(b"contenido"), "historial.png"),
                        "modifiedTimes": json.dumps([None]),
                    },
                    content_type="multipart/form-data",
                )
            self.assertEqual(unsupported_response.status_code, 400)

    def test_historical_document_parser_supports_pdf_docx_and_txt(self):
        text = "Operación de inventario y organización del almacén."
        txt = cv_app.extract_historical_document("historial.txt", text.encode())
        self.assertEqual(txt, [text])

        word = Document()
        word.add_paragraph(text)
        word_buffer = io.BytesIO()
        word.save(word_buffer)
        docx = cv_app.extract_historical_document(
            "historial.docx", word_buffer.getvalue()
        )
        self.assertIn(text, docx)

        pdf_document = fitz.open()
        pdf_document.new_page().insert_text((72, 72), text)
        pdf_bytes = pdf_document.tobytes()
        pdf_document.close()
        pdf = cv_app.extract_historical_document("historial.pdf", pdf_bytes)
        self.assertIn(text, pdf[0])

        scanned_document = fitz.open()
        scanned_document.new_page()
        scanned_pdf = scanned_document.tobytes()
        scanned_document.close()
        with self.assertRaises(cv_app.AppError):
            cv_app.extract_historical_document("escaneado.pdf", scanned_pdf)

    def test_google_authentication_checks_the_allowlisted_identity(self):
        with (
            patch.dict(
                "os.environ",
                {
                    "CV_ALLOWED_EMAILS": "owner@example.com",
                    "CV_AUTH_REQUIRED": "1",
                },
            ),
            patch(
                "firebase_admin.auth.verify_id_token",
                return_value={
                    "uid": "owner-uid",
                    "auth_time": __import__('session_policy').REAUTHENTICATE_AFTER,
                    "email": "owner@example.com",
                    "email_verified": True,
                },
            ) as verify_token,
            cv_app.app.test_request_context(
                "/api/profile-summary",
                headers={"Authorization": "Bearer test-token"},
            ),
        ):
            cv_app.verify_google_identity()

        verify_token.assert_called_once_with("test-token", check_revoked=True)
        with (
            patch.dict("os.environ", {"CV_ALLOWED_EMAILS": "owner@example.com"}),
            patch(
                "firebase_admin.auth.verify_id_token",
                return_value={
                    "email": "other@example.com",
                    "email_verified": True,
                },
            ),
            cv_app.app.test_request_context(
                "/api/profile-summary",
                headers={"Authorization": "Bearer test-token"},
            ),
        ):
            with self.assertRaises(cv_app.AppError) as rejected:
                cv_app.verify_google_identity()
        self.assertEqual(rejected.exception.status_code, 403)


if __name__ == "__main__":
    unittest.main()
