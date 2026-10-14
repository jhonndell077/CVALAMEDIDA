import copy
import io
import unittest
from unittest.mock import patch
import pymupdf
from docx import Document
import app
from drive_bank import clean_facts
from test_drive_bank import memory, MemoryBank
import telegram_bot as bot
from test_telegram import MemoryStore, CONFIG, message


class ProjectCVTests(unittest.TestCase):
    def analysis(self):
        return {"vacancy": {"position": "Especialista en desarrollo web, IA y automatización", "mandatory_requirements": ["HTML, JavaScript, Firebase y aplicaciones web"]}, "matching": {"selected_experiences": [], "selected_skills": [], "selected_project_ids": ["portfolio"], "items": [{"requirement": "Aplicaciones web", "project_ids": ["portfolio", "inventado"], "type": "directo"}]}}

    def test_projects_only_are_valid_and_appear_in_both_exports(self):
        profile = copy.deepcopy(app.MASTER_PROFILE)
        profile["projects"] = [{"id": "portfolio", "title": "Mi aplicación web", "description": "Aplicación web para gestionar inventario con Firebase.", "technologies": ["JavaScript", "Firebase"], "url": "https://example.com/portfolio"}]
        with patch.object(app, "MASTER_PROFILE", profile):
            analysis = app.sanitize_analysis(self.analysis())
            self.assertEqual(analysis["matching"]["items"][0]["project_ids"], ["portfolio"])
            draft = app.make_draft(analysis)
            self.assertTrue(draft["project_focus"])
            self.assertEqual(draft["experiences"], [])
            self.assertIn("Experiencia práctica", draft["summary"])
            self.assertNotIn("Asistente Administrativo", draft["summary"])
            pdf = pymupdf.open(stream=app.make_pdf(draft).getvalue(), filetype="pdf")
            self.assertIn("Mi aplicación web", "".join(page.get_text() for page in pdf))
            document = Document(app.make_docx(draft))
            text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn("Mi aplicación web", text)
            self.assertNotIn("https://example.com/portfolio", text)

    def test_private_document_links_and_provenance_never_leave_cv_engine(self):
        profile = copy.deepcopy(app.MASTER_PROFILE)
        private_url = "https://drive.google.com/drive/folders/private-bank"
        profile["projects"] = [{"id":"portfolio", "title":"Mi aplicación web", "description":"Aplicación web. " + private_url, "technologies":["JavaScript"], "url":private_url, "source_documents":[{"file_id":"private-file", "name":"private-source"}]}]
        with patch.object(app, 'MASTER_PROFILE', profile):
            draft = app.make_draft(self.analysis())
        self.assertNotIn('drive.google.com', str(draft))
        self.assertNotIn('private-file', str(draft))
        # Exporters independently defend against older drafts or direct callers.
        draft['summary'] += ' ' + private_url
        draft['projects'][0]['url'] = private_url
        pdf = pymupdf.open(stream=app.make_pdf(draft).getvalue(), filetype='pdf')
        self.assertNotIn('drive.google.com', ''.join(page.get_text() for page in pdf))
        document = Document(app.make_docx(draft))
        self.assertNotIn('drive.google.com', '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells))

    def test_unknown_project_cannot_establish_experience(self):
        analysis = self.analysis()
        analysis["matching"]["selected_project_ids"] = ["inventado"]
        with self.assertRaises(app.AppError):
            app.sanitize_analysis(analysis)

    def test_telegram_uses_project_only_cv_engine_and_delivers_pdf(self):
        profile = copy.deepcopy(app.MASTER_PROFILE)
        profile["projects"] = [{"id": "portfolio", "title": "Mi aplicación web", "description": "Gestión de inventario web con Firebase.", "technologies": ["JavaScript"], "url": ""}]
        store = MemoryStore()
        store.write('1', {'state': 'queued', 'message': message()}, 0)
        captured = []
        def archive(core, config, identifier, filename, content):
            captured.append(content)
            return 'https://drive.google.com/file/d/test/view'
        with patch.object(app, 'MASTER_PROFILE', profile), patch.object(app, 'load_drive_bank', return_value={'pending':0}), patch.object(app, 'analyze_with_model', return_value=self.analysis()), patch.object(bot, 'drive_token', return_value='token'), patch.object(bot, 'say'), patch.object(bot, 'vacancy_payload', return_value={'text':'vacante'}), patch.object(bot, 'archive_pdf', side_effect=archive), patch.object(bot, 'telegram', return_value={'message_id':42}):
            bot.process_job(app, CONFIG, '1', store)
        self.assertEqual(store.read('1')[0]['state'], 'done')
        document = pymupdf.open(stream=captured[0], filetype='pdf')
        self.assertIn('Mi aplicación web', ''.join(page.get_text() for page in document))

    def test_description_preserves_complete_sentences(self):
        self.assertEqual(app.brief_project_description('Aplicación web. Integración de herramientas y servicios adicionales.', 30), 'Aplicación web.')

    def test_project_folder_is_one_record_and_source_changes_force_reread(self):
        files, contents = memory()
        del files["certificate"]
        files["project"] = {"id": "project", "name": "Portfolio", "mimeType": "application/vnd.google-apps.folder", "parents": ["bank"]}
        files["code"] = {"id": "code", "name": "app.js", "mimeType": "text/javascript", "parents": ["project"], "md5Checksum": "v1"}
        contents["code"] = b'const client_secret = "private";\nfunction inventory() {}'
        calls = []
        def extract(*args):
            calls.append(args)
            return clean_facts({"projects": [{"title": "Portfolio", "description": "Gestión de inventario web", "technologies": ["JavaScript"]}]})
        bank = MemoryBank(files, contents, extract)
        _, status, _ = bank.load()
        self.assertEqual(status["pending"], 1)
        profile, status, _ = bank.load(process=True)
        self.assertEqual(status["pending"], 0)
        self.assertNotIn(b'private', calls[0][2])
        self.assertEqual(profile["projects"][-1]["url"], "https://drive.google.com/drive/folders/project")
        bank.load(process=True)
        self.assertEqual(len(calls), 1)
        files["code"]["md5Checksum"] = "v2"
        self.assertEqual(bank.load()[1]["pending"], 1)
        bank.load(process=True)
        self.assertEqual(len(calls), 2)
